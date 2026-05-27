"""Export the entire articles table to an .xlsx file.

This is the RAW firehose — every article every source has fetched,
regardless of whether Gemini decided it was a homicide sentencing,
whether the sentencing date was recent enough, etc. The extraction_status
column tells you what happened to each one:

  - extracted: passed all filters, became (or matched) a Case
  - no_match:  Gemini said not a homicide sentencing, OR sentencing date
               was too old, OR no defendant name extractable
  - failed:    Gemini API call failed (transient error, JSON parse, etc.)
  - pending:   in flight when the export ran

Run:
    python export_articles.py [output_path]

Default output: ./articles_export_<UTC-timestamp>.xlsx in cwd.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import psycopg2
import psycopg2.extras
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Use the same DATABASE_URL that the rest of the backend reads from .env.
# psycopg2 doesn't understand the SQLAlchemy-style "postgresql+psycopg2://"
# prefix, so strip the dialect tag before handing it to psycopg2.connect().
from app.config import settings

DB_URL = settings.database_url.replace("postgresql+psycopg2://", "postgresql://", 1)


def _flatten(value) -> str:
    """Render any JSON/SQL value as a string suitable for an Excel cell."""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        try:
            return json.dumps(value, ensure_ascii=False, default=str)
        except Exception:
            return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)


def export(output_path: Path) -> Path:
    sql = """
        SELECT
            a.id              AS article_id,
            a.source_name,
            a.source_type,
            a.url,
            a.title,
            a.published_at,
            a.extraction_status,
            a.extraction_model,
            a.extraction_error,
            a.case_id,
            a.created_at,
            a.raw_text,
            a.extracted_json,
            c.defendant_name  AS case_defendant_name,
            c.defendant_age   AS case_defendant_age,
            c.state           AS case_state,
            c.county          AS case_county,
            c.court_name      AS case_court_name,
            c.judge_name      AS case_judge_name,
            c.docket_number   AS case_docket_number,
            c.sentencing_date AS case_sentencing_date,
            c.sentence_text   AS case_sentence_text,
            c.sentence_type   AS case_sentence_type,
            c.sentence_years  AS case_sentence_years,
            c.summary         AS case_summary,
            c.status          AS case_status,
            c.content_score   AS case_content_score
        FROM articles a
        LEFT JOIN cases c ON a.case_id = c.id
        ORDER BY a.created_at DESC
    """

    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cur.execute(sql)
    rows = cur.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "All Articles"

    if not rows:
        ws["A1"] = "No articles in DB yet."
        wb.save(output_path)
        return output_path

    headers = list(rows[0].keys())

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F2937")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, header in enumerate(headers, start=1):
            ws.cell(row=r_idx, column=c_idx, value=_flatten(row[header]))

    # Reasonable column widths
    width_map = {
        "article_id": 38, "source_name": 22, "source_type": 14,
        "url": 60, "title": 60, "published_at": 22,
        "extraction_status": 16, "extraction_model": 16, "extraction_error": 30,
        "case_id": 38, "created_at": 22,
        "raw_text": 80, "extracted_json": 80,
        "case_defendant_name": 24, "case_defendant_age": 10,
        "case_state": 8, "case_county": 18, "case_court_name": 28,
        "case_judge_name": 22, "case_docket_number": 18,
        "case_sentencing_date": 14, "case_sentence_text": 30,
        "case_sentence_type": 14, "case_sentence_years": 12,
        "case_summary": 80, "case_status": 12, "case_content_score": 12,
    }
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width_map.get(header, 20)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) > 1:
        out = Path(sys.argv[1])
    else:
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        out = Path.cwd() / f"articles_export_{ts}.xlsx"

    print(f"Exporting articles table to: {out}")
    path = export(out)

    # Tiny summary
    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM articles")
    total = cur.fetchone()[0]
    cur.execute("""
        SELECT extraction_status, COUNT(*)
        FROM articles
        GROUP BY extraction_status
        ORDER BY COUNT(*) DESC
    """)
    by_status = cur.fetchall()
    cur.execute("""
        SELECT source_name, COUNT(*)
        FROM articles
        GROUP BY source_name
        ORDER BY COUNT(*) DESC
    """)
    by_source = cur.fetchall()
    conn.close()

    print(f"\nTOTAL ARTICLES: {total}")
    print("\nBy extraction_status:")
    for s, n in by_status:
        print(f"  {s:<12} {n}")
    print("\nBy source_name (top 15):")
    for s, n in by_source[:15]:
        print(f"  {s:<30} {n}")
    print(f"\nFile saved to: {path}")
